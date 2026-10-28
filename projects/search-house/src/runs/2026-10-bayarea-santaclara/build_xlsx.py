import os, sys
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, "..", ".."))  # src/, for fonts.py
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "output", "BayArea_House_Shortlist.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
from data import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E78")
N = Font(name=F); B = Font(name=F, bold=True); BLUE = Font(name=F, color="0000FF")
LINK = Font(name=F, color="0563C1", underline="single"); Y = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF"); BOR = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

def header(ws, cols, widths):
    for c, (h, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = H, HF, Alignment(wrap_text=True, vertical="center"), BOR
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "B2"

wb = Workbook()
ws = wb.active; ws.title = "Read me"
ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 100
rows = [
 ("House search: Santa Clara, Sunnyvale, Mountain View, North San Jose", None), ("Prepared", "27 Oct 2026"), (None, None),
 ("Budget (USD)", BUDGET), ("Stretch limit (USD, +5%)", "=B4*1.05"), ("Minimum living area (sq ft)", MINSQFT), (None, None),
 ("Request", "Single-family detached house, 3+ bedrooms, 2+ baths, garage; commute to " + WORK + " under 20 minutes."),
 ("Sheets", "Websites = where to search. All listings = every listing reviewed with size/budget checks. Shortlist = homes researched by agents: grocery, school rating, hospital, safety & hazards, commute, scores."),
 ("Edit", "Yellow cells are inputs: change the budget or minimum size and every pass/fail check recalculates. Blue cells are prices and scores you can update."),
 ("Scoring", "Safety, Livability, Value, Verification are 1-5 from the research agents. Total = Safety + Livability + Value + 2 x Verification. The main pick also had to meet the 20-minute commute."),
 ("Ongoing costs", "California property tax is roughly 1.1-1.25% of the purchase price per year (~$22-25K on $2M), plus homeowner's insurance (earthquake cover is separate) and any HOA / Mello-Roos."),
 ("Caveats", "Zillow, Redfin and homes.com detail pages often blocked automated reading, so some facts come from MLS/brokerage pages or same-plan neighbours. Distances and commute times are estimates; check FEMA flood zones at msc.fema.gov and school assignment by address."),
]
for i, (a, b) in enumerate(rows, 1):
    ca = ws.cell(row=i, column=1, value=a); ca.font = B if i > 1 else Font(name=F, bold=True, size=14)
    cb = ws.cell(row=i, column=2, value=b); cb.font = N; cb.alignment = WRAP
for r in (4, 6):
    ws.cell(row=r, column=2).font = BLUE; ws.cell(row=r, column=2).fill = Y
ws["B4"].number_format = "$#,##0"; ws["B5"].number_format = "$#,##0"; ws["B6"].number_format = "#,##0"
BUD, STR, MIN = "'Read me'!$B$4", "'Read me'!$B$5", "'Read me'!$B$6"

ws = wb.create_sheet("Websites")
header(ws, ["Website", "URL", "Why use it"], [36, 38, 90])
for i, (n, u, w) in enumerate(WEBSITES, 2):
    ws.cell(row=i, column=1, value=n).font = B
    c = ws.cell(row=i, column=2, value=u); c.hyperlink = u; c.font = LINK
    ws.cell(row=i, column=3, value=w).font = N
    for col in range(1, 4): ws.cell(row=i, column=col).alignment = WRAP; ws.cell(row=i, column=col).border = BOR

ws = wb.create_sheet("All listings")
header(ws, ["#", "City", "Address", "Link", "Price ($)", "Beds", "Baths", "Living sq ft", "$ / sq ft", ">= min sq ft?", "3+ bd & 2+ ba?", "Within budget?", "Outcome / reason"],
       [5, 15, 36, 10, 13, 6, 6, 10, 9, 10, 10, 13, 40])
for i, a in enumerate(ALL, 2):
    city, addr, url, price, bd, ba, sq, note = a
    vals = [i - 1, city, addr, "Open listing", price, bd, ba, sq, f"=E{i}/H{i}", f'=IF(H{i}>={MIN},"Yes","No")',
            f'=IF(AND(F{i}>=3,G{i}>=2),"Yes","No")', f'=IF(E{i}<={BUD},"Yes",IF(E{i}<={STR},"Stretch","No"))', note]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=c, value=v); cell.font = N; cell.border = BOR; cell.alignment = WRAP
    ws.cell(row=i, column=4).hyperlink = url; ws.cell(row=i, column=4).font = LINK
    ws.cell(row=i, column=5).font = BLUE; ws.cell(row=i, column=5).number_format = "$#,##0"
    ws.cell(row=i, column=8).number_format = "#,##0"; ws.cell(row=i, column=9).number_format = "$#,##0"
ws.auto_filter.ref = f"A1:M{len(ALL)+1}"

ws = wb.create_sheet("Shortlist")
cols = ["ID", "City", "Address", "Link", "Living sq ft", "Lot (sq ft)", "Bed/Bath", "Garage", "Year", "HOA", "Status / history", "Price ($)", "$ / sq ft",
        "Grocery", "Elementary school (rating)", "Hospital (ER)", "Safety & hazards", "Commute to Marvell", "Red flags",
        "Safety 1-5", "Livability 1-5", "Value 1-5", "Verification 1-5", "Total score", "Rank (all)", "Within budget?", "Status"]
header(ws, cols, [7, 13, 30, 9, 8, 10, 11, 14, 6, 9, 22, 12, 8, 24, 28, 24, 32, 12, 34, 7, 8, 7, 9, 7, 7, 9, 18])
n = len(SL)
for i, r in enumerate(SL, 2):
    (id_, city, addr, sq, lot, bdba, gar, yr, hoa, status, price, url, gro, sch, hos, safe, com, flags, S, L, V, Ver, avoid) = r
    st = "MAIN PICK" if id_ == PICK else ("Avoid" if avoid else (f"Alternative #{ALT.index(id_)+1}" if id_ in ALT else ""))
    vals = [id_, city, addr, "Open listing", sq, lot, bdba, gar, yr, hoa, status, price, f"=L{i}/E{i}", gro, sch, hos, safe, com, flags,
            S, L, V, Ver, f"=T{i}+U{i}+V{i}+2*W{i}", f"=RANK(X{i},$X$2:$X${n+1})", f'=IF(L{i}<={BUD},"Yes",IF(L{i}<={STR},"Stretch","No"))', st]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=c, value=v); cell.font = N; cell.border = BOR; cell.alignment = WRAP
    ws.cell(row=i, column=4).hyperlink = url; ws.cell(row=i, column=4).font = LINK
    for c in (12, 20, 21, 22, 23): ws.cell(row=i, column=c).font = BLUE
    ws.cell(row=i, column=12).number_format = "$#,##0"; ws.cell(row=i, column=13).number_format = "$#,##0"; ws.cell(row=i, column=5).number_format = "#,##0"
    ws.cell(row=i, column=24).font = B
    fill = "E2EFDA" if id_ == PICK else ("FCE4D6" if avoid else None)
    if fill:
        for c in range(1, 28): ws.cell(row=i, column=c).fill = PatternFill("solid", fgColor=fill)
ws.auto_filter.ref = f"A1:AA{n+1}"
ws.cell(row=n + 3, column=1, value="Green = main pick; orange = avoid (townhouse or no garage). Distances and commute times are estimates. Blue cells are inputs.").font = Font(name=F, italic=True)
ws["W1"].comment = Comment("1-5: how much of the listing (size, garage, status, history) the agent could confirm", "Claude")
wb.save(OUT); print("saved")
