import os, sys
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(HERE, "..", ".."))  # src/, for fonts.py
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "output", "HCMC_House_Shortlist.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)
from data import *
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

for i, r in enumerate(SL):
    assert r[10] == next(a[2] for a in ALL if a[8].startswith("Shortlisted " + r[0]))

F = "Arial"
H = Font(name=F, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="1F4E78")
N = Font(name=F)
B = Font(name=F, bold=True)
BLUE = Font(name=F, color="0000FF")
LINK = Font(name=F, color="0563C1", underline="single")
Y = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF")
BOR = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

def header(ws, row, cols, widths):
    for c, (h, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = H, HF, Alignment(wrap_text=True, vertical="center"), BOR
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=2)

wb = Workbook()
# --- Read me / assumptions
ws = wb.active
ws.title = "Read me"
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 95
rows = [
    ("House search: District 10 & Thủ Đức, HCMC", None),
    ("Prepared", "27 Oct 2026"),
    (None, None),
    ("Exchange rate (VND per USD)", RATE),
    ("Budget (USD)", BUDGET),
    ("Budget in VND billion (tỷ)", "=B5*B4/1E9"),
    ("Minimum land area (m2)", 50),
    (None, None),
    ("Sheets", "Websites = best sites to search. All listings = every ad collected, with budget/area checks. Shortlist = houses passing the filter, with nearby market, supermarket, primary school, hospital, safety notes and scores."),
    ("Edit", "Yellow cells are inputs. Change the exchange rate, budget or minimum area and every USD price and pass/fail check recalculates."),
    ("Scoring", "Safety, Livability, Value and Verification are 1-5 from the research agents. Total = Safety + Livability + Value + 2 x Verification (title/address confirmation counts double, because unverified cheap listings are the main risk in this market)."),
    ("Ward names", "HCMC merged wards on 1 July 2025: old D10 wards are now Vườn Lài, Hòa Hưng, Diên Hồng; old Thủ Đức wards are now Hiệp Bình, Tam Bình, Linh Xuân, Thủ Đức, etc."),
    ("Caveats", "Distances are estimates from street geography, not map measurements - check on Google Maps. Several listing pages blocked automated access, so some details come from site index pages. Always see the sổ hồng (title), hoàn công (completion certificate) and a planning (quy hoạch) check before any deposit."),
    ("Rate source", "Vietcombank transfer buying rate 25,660 on 27/10/2026: https://webgia.com/ty-gia/vietcombank/"),
]
for i, (a, b) in enumerate(rows, 1):
    ca = ws.cell(row=i, column=1, value=a); ca.font = B if i > 1 else Font(name=F, bold=True, size=14)
    cb = ws.cell(row=i, column=2, value=b); cb.font = N; cb.alignment = WRAP
for r in (4, 5, 7):
    ws.cell(row=r, column=2).font = BLUE; ws.cell(row=r, column=2).fill = Y
ws["B4"].number_format = "#,##0"; ws["B5"].number_format = "$#,##0"; ws["B6"].number_format = "0.00"
RATE_REF, BUD_REF, AREA_REF = "'Read me'!$B$4", "'Read me'!$B$5", "'Read me'!$B$7"

# --- Websites
ws = wb.create_sheet("Websites")
header(ws, 1, ["Website", "URL", "Why use it"], [36, 38, 90])
for i, (n, u, w) in enumerate(WEBSITES, 2):
    ws.cell(row=i, column=1, value=n).font = B
    c = ws.cell(row=i, column=2, value=u); c.hyperlink = u; c.font = LINK
    ws.cell(row=i, column=3, value=w).font = N
    for col in range(1, 4):
        ws.cell(row=i, column=col).alignment = WRAP; ws.cell(row=i, column=col).border = BOR

# --- All listings
ws = wb.create_sheet("All listings")
cols = ["#", "Site", "Listing title", "Link", "Area", "Ward (new/old)", "Land m2", "Price (tỷ VND)", "Posted", "Price (USD)", "tr VND / m2", "> min area?", "Within budget?", "Outcome / reason"]
header(ws, 1, cols, [5, 11, 40, 12, 6, 22, 9, 11, 11, 12, 11, 10, 11, 42])
for i, a in enumerate(ALL, 2):
    site, title, url, zone, ward, area, price, posted, note = a
    vals = [i - 1, site, title, "Open ad", "D10" if zone == "D10" else "Thủ Đức", ward, area, price, posted,
            f"=H{i}*1E9/{RATE_REF}", f"=H{i}*1000/G{i}", f'=IF(G{i}>{AREA_REF},"Yes","No")', f'=IF(J{i}<={BUD_REF},"Yes","No")', note]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=c, value=v); cell.font = N; cell.border = BOR; cell.alignment = WRAP
    ws.cell(row=i, column=4).hyperlink = url; ws.cell(row=i, column=4).font = LINK
    ws.cell(row=i, column=8).font = BLUE
    ws.cell(row=i, column=10).number_format = "$#,##0"
    ws.cell(row=i, column=11).number_format = "0"
    ws.cell(row=i, column=8).number_format = "0.00"
ws.auto_filter.ref = f"A1:N{len(ALL)+1}"

# --- Shortlist
ws = wb.create_sheet("Shortlist")
cols = ["ID", "Area", "Address", "Ad link", "Land m2", "Floors", "Bed/WC", "Alley / car access", "Legal (title)", "Posted",
        "Price (tỷ)", "Price (USD)", "tr VND/m2", "Market", "Supermarket", "Primary school", "Hospital", "Safety & flooding", "Red flags",
        "Safety 1-5", "Livability 1-5", "Value 1-5", "Verification 1-5", "Total score", "Rank (all)", "Status"]
widths = [7, 7, 30, 9, 7, 13, 9, 18, 16, 11, 8, 11, 8, 26, 26, 26, 28, 34, 36, 7, 8, 7, 9, 7, 7, 16]
header(ws, 1, cols, widths)
n = len(SL)
pick = "TD-14"
for i, r in enumerate(SL, 2):
    (id_, zone, addr, land, floors, pnwc, alley, legal, posted, price, url, mkt, sup, sch, hos, safe, flags, S, L, V, Ver, avoid) = r
    status = "MAIN PICK" if id_ == pick else ("Avoid" if avoid else "")
    vals = [id_, "D10" if zone == "D10" else "Thủ Đức", addr, "Open ad", land, floors, pnwc, alley, legal, posted, price,
            f"=K{i}*1E9/{RATE_REF}", f"=K{i}*1000/E{i}", mkt, sup, sch, hos, safe, flags, S, L, V, Ver,
            f"=T{i}+U{i}+V{i}+2*W{i}", f"=RANK(X{i},$X$2:$X${n+1})", status]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=c, value=v); cell.font = N; cell.border = BOR; cell.alignment = WRAP
    ws.cell(row=i, column=4).hyperlink = url; ws.cell(row=i, column=4).font = LINK
    for c in (11, 20, 21, 22, 23):
        ws.cell(row=i, column=c).font = BLUE
    ws.cell(row=i, column=12).number_format = "$#,##0"; ws.cell(row=i, column=13).number_format = "0"
    ws.cell(row=i, column=24).font = B
    if status == "MAIN PICK":
        for c in range(1, 27): ws.cell(row=i, column=c).fill = PatternFill("solid", fgColor="E2EFDA")
    elif avoid:
        for c in range(1, 27): ws.cell(row=i, column=c).fill = PatternFill("solid", fgColor="FCE4D6")
ws.auto_filter.ref = f"A1:Z{n+1}"
ws.cell(row=n + 3, column=1, value="Green = main pick; orange = avoid (stale, duplicate, rooming house or implausible price). Scores from research agents (blue = inputs you can change). Distances are estimates.").font = Font(name=F, italic=True)
ws["T1"].comment = Comment("1-5 from research agents: security reputation, flooding, traffic, planning", "Claude")
ws["W1"].comment = Comment("1-5: how much of the ad (address, title, size) could be confirmed", "Claude")

out = OUT
wb.save(out)
print("saved", out)
